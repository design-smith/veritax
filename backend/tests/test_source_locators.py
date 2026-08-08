from io import BytesIO

import pytest

from app.source_locators import validate_source_quote


def _pdf_bytes(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    objects = [
        "1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n",
        "2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj\n",
        (
            "3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            "/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >> endobj\n"
        ),
        "4 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n",
    ]
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET"
    objects.append(f"5 0 obj << /Length {len(stream)} >> stream\n{stream}\nendstream endobj\n")
    content = "%PDF-1.4\n"
    offsets = [0]
    for obj in objects:
        offsets.append(len(content.encode()))
        content += obj
    xref_at = len(content.encode())
    content += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n"
    content += "".join(f"{offset:010d} 00000 n \n" for offset in offsets[1:])
    content += f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n"
    return content.encode()


def _docx_bytes(paragraphs: list[str]) -> bytes:
    from docx import Document

    doc = Document()
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def _xlsx_bytes(rows: list[list[str]], *, sheet_name: str = "Sheet1") -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def test_pdf_quote_requires_page_when_text_layer_exists():
    data = _pdf_bytes("The cost plus five percent markup applies.")

    with pytest.raises(ValueError, match="page is required"):
        validate_source_quote("agreement.pdf", "application/pdf", data, "five percent")

    region = validate_source_quote("agreement.pdf", "application/pdf", data, "five percent", page=1)

    assert region.page == 1
    assert region.locator == "page 1"
    assert "five percent" in region.text


def test_txt_quote_uses_line_locator():
    data = b"Header\nCost plus five percent services fee\nFooter\n"

    region = validate_source_quote("notes.txt", "text/plain", data, "five percent")

    assert region.locator == "line 2"
    assert region.page is None


def test_docx_quote_uses_paragraph_locator():
    data = _docx_bytes(["Intro", "The service fee is cost plus five percent.", "End"])

    region = validate_source_quote(
        "services.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        data,
        "five percent",
    )

    assert region.locator == "paragraph 2"


def test_csv_quote_uses_row_locator():
    data = b"Account,Amount\nServices revenue,1200\nRoyalty expense,300\n"

    region = validate_source_quote("trial-balance.csv", "text/csv", data, "Services revenue")

    assert region.locator == "row 2"


def test_xlsx_quote_uses_sheet_and_row_locator():
    data = _xlsx_bytes(
        [["Account", "Amount"], ["Services revenue", "1200"]],
        sheet_name="Trial Balance",
    )

    region = validate_source_quote(
        "trial-balance.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        data,
        "Services revenue",
    )

    assert region.locator == "Trial Balance!R2"


def test_unverifiable_quote_is_rejected():
    with pytest.raises(ValueError, match="quote not found"):
        validate_source_quote("notes.txt", "text/plain", b"Only actual text", "invented quote")
